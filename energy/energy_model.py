"""
储能源荷储匹配分析——纯计算模块（无 Streamlit 依赖）。
从 energy_app.py 中抽出，供独立版 (energy_app.py) 与综合版 (integrated_app.py) 共用。
"""

import io
import math

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font


def generate_8760_month_array():
    days_in_month = [31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    hours_in_month = [d * 24 for d in days_in_month]
    month_array = np.zeros(8760, dtype=int)
    current_hour_index = 0
    for month in range(1, 13):
        num_hours = hours_in_month[month - 1]
        if current_hour_index + num_hours > 8760:
            num_hours = 8760 - current_hour_index
        month_array[current_hour_index:current_hour_index + num_hours] = month
        current_hour_index += num_hours
    return month_array


def calculate_single_case(
    pv_unit_data, wind_unit_data, load_data,
    pv_capacity_1mw, wind_capacity_1mw,
    storage_power_mw, storage_duration_h,
    storage_efficiency_p3, discharge_depth_p2,
    peak_valley_map, prices, month_8760_array,
    discharge_allowed, return_hourly=False,
    hourly_self_prices=None, hourly_on_grid_prices=None
):
    pv_generation = pv_unit_data * pv_capacity_1mw
    wind_generation = wind_unit_data * wind_capacity_1mw
    generation_data = pv_generation + wind_generation

    storage_capacity_max_kwh = storage_power_mw * storage_duration_h * 1000
    max_charge_power_kwh = storage_power_mw * 1000
    max_discharge_power_kwh = storage_power_mw * 1000
    storage_efficiency_single = math.sqrt(storage_efficiency_p3)
    storage_capacity_arr = np.zeros(8760)

    total_consumption_sum = 0.0
    total_on_grid_sum = 0.0
    total_charge_loss = 0.0
    total_discharge_loss = 0.0
    total_discharge_energy = 0.0

    time_period_stats = {
        '尖峰': {'consumption': 0.0, 'on_grid': 0.0, 'consumption_cost_sum': 0.0, 'on_grid_cost_sum': 0.0},
        '峰':   {'consumption': 0.0, 'on_grid': 0.0, 'consumption_cost_sum': 0.0, 'on_grid_cost_sum': 0.0},
        '平':   {'consumption': 0.0, 'on_grid': 0.0, 'consumption_cost_sum': 0.0, 'on_grid_cost_sum': 0.0},
        '谷':   {'consumption': 0.0, 'on_grid': 0.0, 'consumption_cost_sum': 0.0, 'on_grid_cost_sum': 0.0},
        '深谷': {'consumption': 0.0, 'on_grid': 0.0, 'consumption_cost_sum': 0.0, 'on_grid_cost_sum': 0.0}
    }

    if return_hourly:
        hourly_pv_gen = np.zeros(8760)
        hourly_wind_gen = np.zeros(8760)
        hourly_generation = np.zeros(8760)
        hourly_load = np.zeros(8760)
        hourly_charge_in = np.zeros(8760)
        hourly_charge_source = np.zeros(8760)
        hourly_discharge_req = np.zeros(8760)
        hourly_discharge_out = np.zeros(8760)
        hourly_soc = np.zeros(8760)
        hourly_consumption = np.zeros(8760)
        hourly_on_grid = np.zeros(8760)
        hourly_period = ['平'] * 8760
        hourly_charge_loss = np.zeros(8760)
        hourly_discharge_loss = np.zeros(8760)

    use_hourly_prices = (hourly_self_prices is not None) and (hourly_on_grid_prices is not None)

    for i in range(8760):
        current_generation = generation_data[i]
        current_load = load_data[i]
        previous_storage_capacity = storage_capacity_arr[i - 1] if i > 0 else 0.0
        hour_of_day = i % 24
        month = month_8760_array[i]
        current_period_type = peak_valley_map.get(f"{hour_of_day}_{month}", '平')

        # 充电
        storage_charge_in_i = 0.0
        storage_charge_source = 0.0
        on_grid_i = 0.0
        if current_generation > current_load:
            available_from_generation = current_generation - current_load
            remaining_storage_capacity = storage_capacity_max_kwh - previous_storage_capacity
            charge_effective = available_from_generation * storage_efficiency_single
            max_storage_charge = max_charge_power_kwh * storage_efficiency_single
            storage_charge_in_i = min(charge_effective, remaining_storage_capacity,
                                      max_storage_charge / storage_efficiency_single)
            storage_charge_in_i = max(0, storage_charge_in_i)
            storage_charge_source = storage_charge_in_i / storage_efficiency_single
            total_charge_loss += (storage_charge_source - storage_charge_in_i)
            on_grid_i = max(0, available_from_generation - storage_charge_source)

        # 放电（自动检测小时制/时段制）
        storage_required_discharge_i = 0.0
        storage_discharge_out_i = 0.0
        _first_dk = next(iter(discharge_allowed), None) if discharge_allowed else None
        _is_hourly = isinstance(_first_dk, int) if _first_dk is not None else False
        allow_discharge = discharge_allowed.get(hour_of_day, False) if _is_hourly else discharge_allowed.get(current_period_type, False)
        if current_load > current_generation and allow_discharge:
            load_gap = current_load - current_generation
            min_storage_capacity = storage_capacity_max_kwh * (1 - discharge_depth_p2)
            max_discharge_from_current_storage = max(0, previous_storage_capacity - min_storage_capacity)
            required_discharge_for_load = load_gap / storage_efficiency_single
            storage_required_discharge_i = min(max_discharge_from_current_storage,
                                               required_discharge_for_load,
                                               max_discharge_power_kwh)
            storage_required_discharge_i = max(0, storage_required_discharge_i)
            storage_discharge_out_i = storage_required_discharge_i * storage_efficiency_single
            total_discharge_loss += (storage_required_discharge_i - storage_discharge_out_i)
            total_discharge_energy += storage_required_discharge_i

        storage_capacity_i = previous_storage_capacity + storage_charge_in_i - storage_required_discharge_i
        storage_capacity_arr[i] = max(0, min(storage_capacity_i, storage_capacity_max_kwh))

        total_available = current_generation + storage_discharge_out_i
        consumption_i = min(total_available, current_load)
        total_consumption_sum += consumption_i

        if current_generation <= current_load and total_available > current_load:
            on_grid_i += (total_available - current_load)
        total_on_grid_sum += on_grid_i

        if return_hourly:
            hourly_pv_gen[i] = pv_generation[i]
            hourly_wind_gen[i] = wind_generation[i]
            hourly_generation[i] = current_generation
            hourly_load[i] = current_load
            hourly_charge_in[i] = storage_charge_in_i
            hourly_charge_source[i] = storage_charge_source
            hourly_discharge_req[i] = storage_required_discharge_i
            hourly_discharge_out[i] = storage_discharge_out_i
            hourly_soc[i] = storage_capacity_arr[i]
            hourly_consumption[i] = consumption_i
            hourly_on_grid[i] = on_grid_i
            hourly_period[i] = current_period_type
            hourly_charge_loss[i] = storage_charge_source - storage_charge_in_i
            hourly_discharge_loss[i] = storage_required_discharge_i - storage_discharge_out_i

        if use_hourly_prices:
            self_price_i = float(hourly_self_prices[i])
            on_grid_price_i = float(hourly_on_grid_prices[i])
        else:
            self_price_i = prices[current_period_type]['self']
            on_grid_price_i = prices[current_period_type]['on_grid']

        if current_period_type in time_period_stats:
            s = time_period_stats[current_period_type]
            s['consumption'] += consumption_i
            s['on_grid'] += on_grid_i
            s['consumption_cost_sum'] += consumption_i * self_price_i
            s['on_grid_cost_sum'] += on_grid_i * on_grid_price_i

    total_pv_gen = np.sum(pv_generation)
    total_wind_gen = np.sum(wind_generation)
    total_curtailment_sum = total_charge_loss + total_discharge_loss + storage_capacity_arr[-1]
    total_generation_sum = total_pv_gen + total_wind_gen
    total_consumption_cost = sum(s['consumption_cost_sum'] for s in time_period_stats.values())
    total_on_grid_cost = sum(s['on_grid_cost_sum'] for s in time_period_stats.values())
    weighted_self_price = total_consumption_cost / total_consumption_sum if total_consumption_sum > 0 else 0.0
    weighted_on_grid_price = total_on_grid_cost / total_on_grid_sum if total_on_grid_sum > 0 else 0.0
    total_revenue = total_consumption_cost + total_on_grid_cost + (total_curtailment_sum * prices['Curtailment'])
    integrated_price = total_revenue / total_generation_sum if total_generation_sum > 0 else 0.0
    equivalent_cycles = total_discharge_energy / storage_capacity_max_kwh if storage_capacity_max_kwh > 0 else 0.0
    pv_hours = total_pv_gen / (pv_capacity_1mw * 1000) if pv_capacity_1mw > 0 else 0.0
    wind_hours = total_wind_gen / (wind_capacity_1mw * 1000) if wind_capacity_1mw > 0 else 0.0

    result = {
        "total_generation_sum": total_generation_sum,
        "total_pv_generation": total_pv_gen,
        "total_wind_generation": total_wind_gen,
        "pv_hours": pv_hours,
        "wind_hours": wind_hours,
        "total_consumption_sum": total_consumption_sum,
        "total_on_grid_sum": total_on_grid_sum,
        "total_curtailment_sum": total_curtailment_sum,
        "weighted_self_price": weighted_self_price,
        "weighted_on_grid_price": weighted_on_grid_price,
        "integrated_price": integrated_price,
        "storage_equivalent_cycles": equivalent_cycles,
        "time_period_stats": time_period_stats
    }
    if return_hourly:
        result["hourly_data"] = {
            "hour": np.arange(8760),
            "month": month_8760_array,
            "hour_of_day": np.arange(8760) % 24,
            "period_type": hourly_period,
            "pv_generation_kwh": hourly_pv_gen,
            "wind_generation_kwh": hourly_wind_gen,
            "total_generation_kwh": hourly_generation,
            "load_kwh": hourly_load,
            "charge_source_kwh": hourly_charge_source,
            "charge_into_storage_kwh": hourly_charge_in,
            "charge_loss_kwh": hourly_charge_loss,
            "discharge_required_kwh": hourly_discharge_req,
            "discharge_out_kwh": hourly_discharge_out,
            "discharge_loss_kwh": hourly_discharge_loss,
            "storage_soc_kwh": hourly_soc,
            "consumption_kwh": hourly_consumption,
            "on_grid_kwh": hourly_on_grid,
        }
    return result


def perform_batch_calculation(pv_unit_data, wind_unit_data, load_data, params, month_8760_array):
    batch_results = []
    total_load_sum = np.sum(load_data)
    for current_pv in params['pv_list']:
        for current_wind in params['wind_list']:
            for current_power in params['power_list']:
                for current_duration in params['duration_list']:
                    res = calculate_single_case(
                        pv_unit_data, wind_unit_data, load_data,
                        current_pv, current_wind, current_power, current_duration,
                        params['efficiency'], params['depth'],
                        params['peak_valley_map'], params['prices'],
                        month_8760_array, params['discharge_allowed'],
                        hourly_self_prices=params.get('hourly_self_prices'),
                        hourly_on_grid_prices=params.get('hourly_on_grid_prices')
                    )
                    batch_results.append({
                        "光伏容量 (MW)": current_pv,
                        "风电容量 (MW)": current_wind,
                        "光伏利用小时数 (h)": res["pv_hours"],
                        "风电利用小时数 (h)": res["wind_hours"],
                        "储能功率 (MW)": current_power,
                        "储能时长 (h)": current_duration,
                        "储能容量 (MWh)": current_power * current_duration,
                        "加权自用电价": res["weighted_self_price"],
                        "加权上网电价": res["weighted_on_grid_price"],
                        "综合电价": res["integrated_price"],
                        "总发电量 (kWh)": res["total_generation_sum"],
                        "总消纳电量 (kWh)": res["total_consumption_sum"],
                        "总上网电量 (kWh)": res["total_on_grid_sum"],
                        "总折损电量 (kWh)": res["total_curtailment_sum"],
                        "自用比例 (%)": (res["total_consumption_sum"] / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "绿电占用电比例 (%)": (res["total_consumption_sum"] / total_load_sum * 100) if total_load_sum > 0 else 0.0,
                        "尖峰消纳 (%)": (res["time_period_stats"]["尖峰"]["consumption"] / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "峰消纳 (%)":   (res["time_period_stats"]["峰"]["consumption"]   / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "平消纳 (%)":   (res["time_period_stats"]["平"]["consumption"]   / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "谷消纳 (%)":   (res["time_period_stats"]["谷"]["consumption"]   / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "深谷消纳 (%)": (res["time_period_stats"]["深谷"]["consumption"] / res["total_generation_sum"] * 100) if res["total_generation_sum"] > 0 else 0.0,
                        "储能等效循环次数": res["storage_equivalent_cycles"]
                    })
    return batch_results


# ===================== Excel 导出 =====================
def write_batch_results_to_excel(results, params):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "批量计算汇总"
    headers = [
        "序号", "光伏容量 (MW)", "风电容量 (MW)", "光伏利用小时数 (h)", "风电利用小时数 (h)",
        "储能功率 (MW)", "储能时长 (h)", "储能容量 (MWh)",
        "加权自用电价", "加权上网电价", "综合电价",
        "总发电量 (kWh)", "消纳总电量 (kWh)", "上网总电量 (kWh)", "折损总电量 (kWh)",
        "自用比例 (%)", "绿电占用电比例 (%)",
        "储能等效循环次数"
    ]
    sheet.append(headers)
    for col_idx, header in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18
    for row_idx, result in enumerate(results, 2):
        sheet.cell(row=row_idx, column=1, value=row_idx - 1)
        sheet.cell(row=row_idx, column=2, value=result['光伏容量 (MW)'])
        sheet.cell(row=row_idx, column=3, value=result['风电容量 (MW)'])
        sheet.cell(row=row_idx, column=4, value=round(result['光伏利用小时数 (h)'], 3))
        sheet.cell(row=row_idx, column=5, value=round(result['风电利用小时数 (h)'], 3))
        sheet.cell(row=row_idx, column=6, value=result['储能功率 (MW)'])
        sheet.cell(row=row_idx, column=7, value=result['储能时长 (h)'])
        sheet.cell(row=row_idx, column=8, value=result['储能容量 (MWh)'])
        sheet.cell(row=row_idx, column=9, value=round(result['加权自用电价'], 4))
        sheet.cell(row=row_idx, column=10, value=round(result['加权上网电价'], 4))
        sheet.cell(row=row_idx, column=11, value=round(result['综合电价'], 4))
        sheet.cell(row=row_idx, column=12, value=round(result['总发电量 (kWh)'], 6))
        sheet.cell(row=row_idx, column=13, value=round(result['总消纳电量 (kWh)'], 6))
        sheet.cell(row=row_idx, column=14, value=round(result['总上网电量 (kWh)'], 6))
        sheet.cell(row=row_idx, column=15, value=round(result['总折损电量 (kWh)'], 6))
        sheet.cell(row=row_idx, column=16, value=round(result['自用比例 (%)'], 2))
        sheet.cell(row=row_idx, column=17, value=round(result['绿电占用电比例 (%)'], 2))
        sheet.cell(row=row_idx, column=18, value=round(result['储能等效循环次数'], 2))
    excel_stream = io.BytesIO()
    workbook.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream


def write_hourly_data_to_excel(hourly_data, scheme_info):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "8760逐时过程量"
    summary_headers = [
        "光伏容量 (MW)", "风电容量 (MW)", "储能功率 (MW)", "储能时长 (h)",
        "储能容量 (MWh)", "综合电价", "总发电量 (kWh)", "总消纳电量 (kWh)",
        "总上网电量 (kWh)", "总折损电量 (kWh)", "自用比例 (%)", "绿电占用电比例 (%)"
    ]
    for col_idx, header in enumerate(summary_headers, 1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="center")
    summary_values = [
        scheme_info.get('光伏容量 (MW)', ''), scheme_info.get('风电容量 (MW)', ''),
        scheme_info.get('储能功率 (MW)', ''), scheme_info.get('储能时长 (h)', ''),
        scheme_info.get('储能容量 (MWh)', ''), round(scheme_info.get('综合电价', 0), 4),
        round(scheme_info.get('总发电量 (kWh)', 0), 2), round(scheme_info.get('总消纳电量 (kWh)', 0), 2),
        round(scheme_info.get('总上网电量 (kWh)', 0), 2), round(scheme_info.get('总折损电量 (kWh)', 0), 2),
        round(scheme_info.get('自用比例 (%)', 0), 2), round(scheme_info.get('绿电占用电比例 (%)', 0), 2),
    ]
    for col_idx, val in enumerate(summary_values, 1):
        cell = sheet.cell(row=2, column=col_idx, value=val)
        cell.alignment = Alignment(horizontal="center")
    start_row = 4
    hourly_headers = [
        "小时序号", "月份", "日内小时", "时段类型",
        "光伏发电 (kWh)", "风电发电 (kWh)", "总发电量 (kWh)", "负载 (kWh)",
        "充电来源 (kWh)", "充入储能 (kWh)", "充电损失 (kWh)",
        "放电需求 (kWh)", "放电输出 (kWh)", "放电损失 (kWh)",
        "储能SOC (kWh)", "消纳量 (kWh)", "上网量 (kWh)"
    ]
    for col_idx, header in enumerate(hourly_headers, 1):
        cell = sheet.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 18
    n = len(hourly_data["hour"])
    for i in range(n):
        row = start_row + 1 + i
        sheet.cell(row=row, column=1, value=i + 1)
        sheet.cell(row=row, column=2, value=int(hourly_data["month"][i]))
        sheet.cell(row=row, column=3, value=int(hourly_data["hour_of_day"][i]))
        sheet.cell(row=row, column=4, value=str(hourly_data["period_type"][i]))
        sheet.cell(row=row, column=5, value=round(float(hourly_data["pv_generation_kwh"][i]), 4))
        sheet.cell(row=row, column=6, value=round(float(hourly_data["wind_generation_kwh"][i]), 4))
        sheet.cell(row=row, column=7, value=round(float(hourly_data["total_generation_kwh"][i]), 4))
        sheet.cell(row=row, column=8, value=round(float(hourly_data["load_kwh"][i]), 4))
        sheet.cell(row=row, column=9, value=round(float(hourly_data["charge_source_kwh"][i]), 4))
        sheet.cell(row=row, column=10, value=round(float(hourly_data["charge_into_storage_kwh"][i]), 4))
        sheet.cell(row=row, column=11, value=round(float(hourly_data["charge_loss_kwh"][i]), 4))
        sheet.cell(row=row, column=12, value=round(float(hourly_data["discharge_required_kwh"][i]), 4))
        sheet.cell(row=row, column=13, value=round(float(hourly_data["discharge_out_kwh"][i]), 4))
        sheet.cell(row=row, column=14, value=round(float(hourly_data["discharge_loss_kwh"][i]), 4))
        sheet.cell(row=row, column=15, value=round(float(hourly_data["storage_soc_kwh"][i]), 4))
        sheet.cell(row=row, column=16, value=round(float(hourly_data["consumption_kwh"][i]), 4))
        sheet.cell(row=row, column=17, value=round(float(hourly_data["on_grid_kwh"][i]), 4))
    sheet.freeze_panes = f"A{start_row + 1}"
    excel_stream = io.BytesIO()
    workbook.save(excel_stream)
    excel_stream.seek(0)
    return excel_stream


# ===================== 输入解析 =====================
def parse_time_slot_input(s):
    try:
        slots = []
        for part in s.replace(' ', '').split(','):
            if not part: continue
            start, end = map(int, part.split('-'))
            slots.append((start, end))
        return slots
    except:
        return None


def parse_batch_input(input_str):
    try:
        parts = [float(x.strip()) for x in input_str.split(',')]
        if len(parts) == 1:
            return [parts[0]]
        elif len(parts) == 3:
            start, end, step = parts
            if step <= 0: return [start]
            return np.arange(start, end + 1e-9, step).tolist()
        else:
            return None
    except:
        return None
